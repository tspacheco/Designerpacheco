"""Gera a base de prospeção do Instagram, o MENSAGENS.md e o gerador HTML.

Uso: python3 gerar.py [contas.json]
Saída (nesta pasta): prospecao-instagram.xlsx, MENSAGENS.md, gerador.html
"""
import json
import os
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import dados as D

AQUI = os.path.dirname(os.path.abspath(__file__))
CAT = {c["id"]: c for c in D.CATEGORIAS}
NOMES_CAT = [c["nome"] for c in D.CATEGORIAS]
ESTADOS = [e[0] for e in D.ESTADOS]
CANAIS = ["DM", "Comentário", "LinkedIn", "E-mail", "Introdução", "Evento"]

COLUNAS = [  # (cabeçalho, largura)
    ("ID", 5), ("Categoria", 22), ("Nome", 28), ("@", 18), ("URL Instagram", 28), ("Cidade", 14),
    ("Língua", 7), ("Seguidores", 12), ("Sinais públicos", 42), ("Evidência (fonte)", 34),
    ("Foco a investigar", 40), ("Razão de encaixe", 34), ("Abordagem recomendada", 36),
    ("Intuito da abordagem", 36), ("Encaixe 0-3", 8), ("Acesso 0-2", 8), ("Complexidade 0-3", 8),
    ("Prova 0-2", 8), ("Canal", 11), ("Gancho", 34), ("Mensagem 1", 60), ("Follow-up", 50),
    ("Estado", 16), ("Data contacto", 12), ("Próxima ação", 28), ("Data próxima ação", 12), ("Notas", 34),
    # Fórmulas no fim: a linha copiada no gerador cola-se de A até Notas sem as apagar.
    ("Score", 7), ("Prioridade", 9), ("Data follow-up", 12), ("Data D7", 12), ("Data D12", 12),
]
COL = {h: i + 1 for i, (h, _) in enumerate(COLUNAS)}


def L(h):
    return get_column_letter(COL[h])


def cabecalho(ws, cols, cor="1F2A44"):
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=cor)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def folha_prospecao(wb, contas):
    ws = wb.active
    ws.title = "Prospeção"
    cabecalho(ws, COLUNAS)
    ws.row_dimensions[1].height = 32
    n_linhas = max(len(contas) + 200, 300)
    for r, c in enumerate(contas, 2):
        cat = CAT[c["categoria"]]
        handle = c.get("instagram", "").lstrip("@")
        val = {
            "ID": r - 1, "Categoria": cat["nome"], "Nome": c["nome"],
            "@": ("@" + handle) if handle else "a confirmar",
            "URL Instagram": ("https://www.instagram.com/" + handle + "/") if handle else "",
            "Cidade": c.get("cidade", ""), "Língua": c.get("lingua", "ro"),
            "Seguidores": c.get("seguidores", "a confirmar"), "Sinais públicos": c.get("sinais", ""),
            "Evidência (fonte)": c.get("evidencia", ""), "Foco a investigar": c.get("foco", ""),
            "Razão de encaixe": c.get("razao", ""), "Abordagem recomendada": cat["primeiro_passo"],
            "Intuito da abordagem": cat["intuito"],
            "Encaixe 0-3": c.get("score_encaixe", 0), "Acesso 0-2": c.get("score_acesso", 0),
            "Complexidade 0-3": c.get("score_complexidade", 0), "Prova 0-2": c.get("score_prova", 0),
            "Canal": "Comentário" if c["categoria"] == "comunidade" else "DM",
            "Gancho": c.get("gancho", ""), "Mensagem 1": c.get("mensagem", ""),
            "Estado": c.get("estado", "Descoberto"),
            "Próxima ação": "Confirmar @ e seguidores no telemóvel" if not handle else "Gerar mensagem e enviar",
            "Data próxima ação": date.today() + timedelta(days=1), "Notas": c.get("notas", ""),
        }
        for h, v in val.items():
            cell = ws.cell(row=r, column=COL[h], value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    # Fórmulas em todas as linhas (incluindo as vazias, para o Tomás acrescentar)
    for r in range(2, n_linhas + 2):
        e, a, co, p = (L(h) + str(r) for h in ("Encaixe 0-3", "Acesso 0-2", "Complexidade 0-3", "Prova 0-2"))
        nome = L("Nome") + str(r)
        ws.cell(row=r, column=COL["Score"], value=f'=IF({nome}="","",{e}+{a}+{co}+{p})')
        s = L("Score") + str(r)
        ws.cell(row=r, column=COL["Prioridade"],
                value=f'=IF({s}="","",IF({s}>=8,5,IF({s}>=6,4,IF({s}>=4,3,IF({s}>=2,2,1)))))')
        dc = L("Data contacto") + str(r)
        cat = L("Categoria") + str(r)
        # Cadência da triagem proativa: D4 observação nova, D7 pergunta sim ou não, D12 fecho limpo
        # (comunidades: 7 dias entre interações, sem D7 nem D12)
        ws.cell(row=r, column=COL["Data follow-up"],
                value=f'=IF({dc}="","",{dc}+IF({cat}="Página / comunidade B2B",7,4))')
        ws.cell(row=r, column=COL["Data D7"],
                value=f'=IF(OR({dc}="",{cat}="Página / comunidade B2B"),"",{dc}+7)')
        ws.cell(row=r, column=COL["Data D12"],
                value=f'=IF(OR({dc}="",{cat}="Página / comunidade B2B"),"",{dc}+12)')
        for h in ("Data contacto", "Data follow-up", "Data D7", "Data D12", "Data próxima ação"):
            ws.cell(row=r, column=COL[h]).number_format = "DD/MM/YYYY"
    fim = n_linhas + 1
    for lista, h in ((NOMES_CAT, "Categoria"), (ESTADOS, "Estado"), (CANAIS, "Canal"), (["ro", "pt", "en"], "Língua")):
        dv = DataValidation(type="list", formula1='"' + ",".join(lista) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{L(h)}2:{L(h)}{fim}")
    for h, mx in (("Encaixe 0-3", 3), ("Acesso 0-2", 2), ("Complexidade 0-3", 3), ("Prova 0-2", 2)):
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2=str(mx), allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{L(h)}2:{L(h)}{fim}")
    rng = f"{L('Prioridade')}2:{L('Prioridade')}{fim}"
    for v, cor in ((5, "1E7B4F"), (4, "6AA84F"), (3, "F1C232"), (2, "E69138")):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[str(v)],
                                      fill=PatternFill("solid", fgColor=cor), font=Font(bold=True)))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{fim}"


def folha_simples(wb, titulo, cols, linhas, cor="1F2A44"):
    ws = wb.create_sheet(titulo)
    cabecalho(ws, cols, cor)
    for r, linha in enumerate(linhas, 2):
        for i, v in enumerate(linha, 1):
            ws.cell(row=r, column=i, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def excel(contas):
    wb = Workbook()
    folha_prospecao(wb, contas)
    folha_simples(wb, "Categorias", [("Categoria", 24), ("Foco", 30), ("Perfil ideal", 44), ("Primeiro passo", 44),
                                     ("Ativo oferecido", 44), ("Intuito", 44), ("Excluir", 44)],
                  [(c["nome"], c["foco"], c["perfil"], c["primeiro_passo"], c["ativo"], c["intuito"], c["excluir"])
                   for c in D.CATEGORIAS])
    linhas = []
    for c in D.CATEGORIAS:
        for lg in ("ro", "pt", "en"):
            for k, t in D.MENSAGENS[c["id"]][lg].items():
                linhas.append((c["nome"], lg, ROTULO[k], t.replace("{fecho}", D.FECHO[lg])))
    folha_simples(wb, "Mensagens", [("Categoria", 22), ("Língua", 7), ("Tipo", 18), ("Modelo", 110)], linhas)
    folha_simples(wb, "Percursos", [("Tipo de negócio", 24), ("Percurso RO", 50), ("Pergunta RO", 50), ("Isca RO", 50),
                                    ("Percurso PT", 50), ("Pergunta PT", 50), ("Isca PT", 50), ("Título do print PT", 34),
                                    ("Passo a validar PT", 44), ("3 validações PT", 60), ("Métrica PT", 44), ("Quick win PT", 44)],
                  [(p["tipo"], p["ro"], p["pergunta"]["ro"], p["isca"]["ro"], p["pt"], p["pergunta"]["pt"], p["isca"]["pt"],
                    p["titulo"]["pt"], p["risco"]["pt"], " / ".join(p["validacoes"]["pt"]), p["metrica"]["pt"], p["quickwin"]["pt"])
                   for p in D.PERCURSOS])
    folha_simples(wb, "Queries", [("Categoria", 24), ("Queries de partida", 100)],
                  list(D.QUERIES) + [("", "")] + [("Termos: " + a, b) for a, b in D.TERMOS])
    folha_simples(wb, "Estados", [("Estado", 24), ("Significado", 70)], D.ESTADOS)
    folha_simples(wb, "Linguagem", [("Usar", 55), ("Evitar", 55)], D.EVITAR)
    saida = os.path.join(AQUI, "prospecao-instagram.xlsx")
    wb.save(saida)
    return saida


ROTULO = {"dm1": "D0: DM 1", "fu": "D4: observação nova", "d7": "D7: pergunta sim ou não", "d12": "D12: fecho limpo",
          "c1": "Comentário: ponto", "c2": "Comentário: acréscimo", "c3": "Comentário: pergunta",
          "triagem": "Triagem (quando respondem, com o Micro System Print)", "parceria": "Parceria (antes de aplicar os 3 quick wins: incentivo primeiro)",
          "testemunho": "Testemunho (no dia em que os 3 quick wins ficam a funcionar)"}


def markdown():
    out = ["# Mensagens de abordagem no Instagram", "",
           "Gerado por `gerar.py` a partir de `dados.py`. Não editar à mão: mudar `dados.py` e correr de novo.", "",
           "Marcadores: `{saudacao}` varia sozinha · `{nome}` · `{empresa}` · `{gancho}` (facto público e concreto da conta) · "
           "`{percurso}` · `{pergunta}` (hipótese em forma de pergunta) · `{isca}` (o que já é possível hoje, em pergunta; vem do tipo de negócio) · "
           "`{tema}` (tema do post) · `{valor}` (observação nova para o D4).", "",
           "**Triagem proativa (playbook 11/10, secção A.14):** a triagem é conduzida por nós. Observar o percurso público, levar uma hipótese "
           "e um Micro System Print de uma página, pedir só a validação com o decisor. Nunca formulário, nunca minutos, nunca afirmar falhas internas. "
           "Cadência: D0 DM → D4 observação nova → D7 pergunta sim ou não → D12 fecho limpo. Quem responde recebe a «Triagem» com o print.", ""]
    for c in D.CATEGORIAS:
        out += ["## " + c["nome"], "", "**Primeiro passo:** " + c["primeiro_passo"], "",
                "**Intuito:** " + c["intuito"], "", "**Registo em romeno:** " + c["registo_ro"], ""]
        for lg, nome in (("ro", "Romeno"), ("pt", "Português"), ("en", "Inglês")):
            out.append("### " + nome)
            out.append("")
            for k, t in D.MENSAGENS[c["id"]][lg].items():
                out += ["**" + ROTULO[k] + "**", "", "> " + t.replace("{fecho}", D.FECHO[lg]), ""]
    out += ["## Micro System Print (1 página)", "",
            "O gerador preenche-o com o tipo de negócio (título, passo a validar, 3 validações, métrica, quick win). "
            "Ajustar ao que se viu: só canais observados e o passo de risco marcado «a validar».", ""]
    for lg, nome in (("ro", "Romeno"), ("pt", "Português"), ("en", "Inglês")):
        out += ["### " + nome, "", "```", D.MICRO_PRINT[lg], "```", ""]
    out += ["## Saudações (rodam sozinhas no gerador)", ""]
    for lg, regs in D.SAUDACOES.items():
        for reg, lst in regs.items():
            out.append("- " + lg.upper() + " · " + reg + ": " + " / ".join(lst))
    out.append("")
    with open(os.path.join(AQUI, "MENSAGENS.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(out))


def html():
    dados = {"categorias": D.CATEGORIAS, "saudacoes": D.SAUDACOES, "fecho": D.FECHO,
             "mensagens": D.MENSAGENS, "percursos": D.PERCURSOS, "evitar": D.EVITAR, "micro": D.MICRO_PRINT}
    with open(os.path.join(AQUI, "gerador.template.html"), encoding="utf-8") as f:
        modelo = f.read()
    js = json.dumps(dados, ensure_ascii=False).replace("</", "<\\/")
    with open(os.path.join(AQUI, "gerador.html"), "w", encoding="utf-8") as f:
        f.write(modelo.replace("/*__DADOS__*/{}", js))


if __name__ == "__main__":
    entrada = sys.argv[1] if len(sys.argv) > 1 else os.path.join(AQUI, "contas.json")
    contas = json.load(open(entrada, encoding="utf-8"))
    print(excel(contas))
    markdown()
    if os.path.exists(os.path.join(AQUI, "gerador.template.html")):
        html()
    print("ok:", len(contas), "contas")
