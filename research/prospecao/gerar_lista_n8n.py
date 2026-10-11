"""Lista para o fluxo do Tomás no n8n (folha «Fila_Envio», mesmas colunas da lista de 50 de 11/10).

Uso: python3 gerar_lista_n8n.py leads.json saida.xlsx primeiro_id
Só entram leads com e-mail confirmado no site oficial (campo email + fonte_email).
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import mensagens as M

COLS = ["id", "cidade_alvo", "localidade_verificada", "empresa", "setor", "email", "website", "assunto", "mensagem",
        "estado", "faturacao_anual_ron", "faturacao_mensal_estimada_ron", "fonte_empresa", "fonte_email",
        "notas_verificacao", "subsetor", "sinais", "percurso", "follow_up_d3", "follow_up_d7", "follow_up_d12"]
NOME_SETOR = {"restauracao": "Restauração / hospitalidade", "marcacao": "Serviços por marcação",
              "construcao": "Construção / serviços por orçamento"}


def main(entrada, saida, primeiro):
    leads = [l for l in json.load(open(entrada, encoding="utf-8")) if l.get("email") and l.get("fonte_email")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Fila_Envio"
    ws.append(COLS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2A44")
    for i, l in enumerate(leads):
        ws.append([
            "romania-2026-%03d" % (primeiro + i), l.get("cidade", "Iași"), l.get("cidade", "Iași"), l["empresa"],
            NOME_SETOR[l["setor"]], l["email"], l["website"], M.assunto_setor(l), M.mensagem_setor(l),
            "Pronto para envio", "a confirmar", "a confirmar", l.get("fonte", ""), l["fonte_email"],
            l.get("notas", ""), l.get("subsetor", ""), "; ".join(l["sinais"]) if isinstance(l.get("sinais"), list) else l.get("sinais", ""), l["percurso_ro"],
            M.cadencia_setor(l, "d3"), M.cadencia_setor(l, "d7"), M.cadencia_setor(l, "d12"),
        ])
    for col, w in zip("ABCDEFGHI", (18, 10, 10, 28, 22, 28, 28, 34, 80)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=c.column_letter == "I", vertical="top")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(saida)
    print("{} leads → {}".format(len(leads), saida))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], int(sys.argv[3]))
