"""Gera a base de prospeção (Excel) a partir de um JSON de contas.

Uso: python3 gerar_base.py contas.json saida.xlsx

O JSON é uma lista de contas (campos em LEIA-ME.md). A base tem e-mails de negócios,
por isso NUNCA vai para o repositório (público): fica em /mnt/project-files/prospecao/.
"""
import json
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import mensagens as M

CATEGORIAS = {"influenciador": "Influenciador / operador", "grande": "Grande negócio",
              "pme_parceiro": "PME / parceiro", "comunidade": "Página / comunidade"}
ESTADOS = ["Descoberto", "Em validação", "Validado", "Mini-mapa em preparação", "Contactado",
           "Resposta", "Reunião marcada", "Não adequado", "Sem resposta"]
CANAIS = ["DM", "E-mail", "LinkedIn", "Comentário", "Introdução", "Evento"]

COLUNAS = [  # (cabeçalho, chave, largura)
    ("ID", None, 6), ("Categoria", None, 20), ("Nome", "nome", 24), ("Instagram", "instagram", 20),
    ("URL Instagram", "url_instagram", 30), ("Cidade", "cidade", 12), ("Site", "site", 26),
    ("LinkedIn", "linkedin", 26), ("E-mail", "email", 26), ("Fonte e-mail", "fonte_email", 26),
    ("Seguidores", "seguidores", 10), ("Língua", "lingua", 7), ("Sinais públicos", "sinais", 40),
    ("Foco a investigar", "foco", 34), ("Razão de encaixe", "razao", 34), ("Evidência", "evidencia", 34),
    ("Decisor", "decisor", 20), ("Encaixe 0-3", "score_encaixe", 8), ("Acesso 0-2", "score_acesso", 8),
    ("Complexidade 0-3", "score_complexidade", 8), ("Prova 0-2", "score_prova", 8),
    ("Score", None, 7), ("Prioridade", None, 9), ("Canal", None, 11), ("Abordagem recomendada", None, 36),
    ("Intuito da abordagem", None, 36), ("Assunto e-mail", None, 30), ("Mensagem", None, 70),
    ("Follow-up", None, 50), ("Link DM", None, 24), ("Estado", None, 14), ("Próxima acção", None, 28),
    ("Data próxima acção", None, 12), ("Data contacto", None, 12), ("Notas", None, 30),
]
COL = {c[0]: i + 1 for i, c in enumerate(COLUNAS)}


def L(nome):
    return get_column_letter(COL[nome])


def preparar(conta, i):
    cat = conta["categoria"]
    canal = M.canal(conta)
    linha = {k: conta.get(k, "") for _, k, _ in COLUNAS if k}
    linha.update({
        "ID": i + 1, "Categoria": CATEGORIAS[cat], "Canal": canal,
        "Abordagem recomendada": M.ABORDAGEM[cat], "Intuito da abordagem": M.INTUITO[cat],
        "Assunto e-mail": M.assunto(conta) if canal == "E-mail" else "",
        "Mensagem": M.email(conta, i) if canal == "E-mail" else M.mensagem(conta, i),
        "Follow-up": "" if cat == "comunidade" else M.followup(conta, i + 1),
        "Link DM": ("https://ig.me/m/" + conta["instagram"].lstrip("@")) if conta.get("instagram") else "",
        "Estado": conta.get("estado") or ("Validado" if conta.get("instagram") or conta.get("email") else "Em validação"),
        "Próxima acção": {"DM": "Enviar DM", "E-mail": "Enviar e-mail (n8n)", "LinkedIn": "Encontrar decisor no LinkedIn",
                          "Comentário": "Comentar o post atual"}.get(canal, "Definir"),
        "Data próxima acção": date.today() + timedelta(days=1),
        "Notas": conta.get("notas", ""),
    })
    for s in ("score_encaixe", "score_acesso", "score_complexidade", "score_prova"):
        linha[s] = int(conta.get(s) or 0)
    return linha


def score_py(c):
    s = sum(int(c.get(k) or 0) for k in ("score_encaixe", "score_acesso", "score_complexidade", "score_prova"))
    return s, (5 if s >= 8 else 4 if s >= 6 else 3 if s >= 4 else 2 if s >= 2 else 1)


def main(entrada, saida):
    contas = json.load(open(entrada, encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospeção"
    cab = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F2A44")
    for i, (h, _, w) in enumerate(COLUNAS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = cab, fundo
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, conta in enumerate(contas, 2):
        linha = preparar(conta, r - 2)
        for h, k, _ in COLUNAS:
            v = linha.get(k) if k else linha.get(h)
            if h == "Score":
                v = "=SUM({0}{2}:{1}{2})".format(L("Encaixe 0-3"), L("Prova 0-2"), r)
            elif h == "Prioridade":
                s = "{}{}".format(L("Score"), r)
                v = "=IF({0}>=8,5,IF({0}>=6,4,IF({0}>=4,3,IF({0}>=2,2,1))))".format(s)
            c = ws.cell(row=r, column=COL[h], value=v)
            c.alignment = Alignment(wrap_text=h in ("Mensagem", "Follow-up", "Sinais públicos", "Foco a investigar",
                                                    "Razão de encaixe", "Evidência"), vertical="top")
            if h.startswith("Data") and v:
                c.number_format = "DD/MM/YYYY"
    n = max(len(contas) + 1, 2)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(COLUNAS)), n)
    for nome, opcoes in (("Categoria", list(CATEGORIAS.values())), ("Estado", ESTADOS), ("Canal", CANAIS)):
        dv = DataValidation(type="list", formula1='"{}"'.format(",".join(opcoes)), allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("{0}2:{0}500".format(L(nome)))

    # Top 5 por categoria (valores, para ler no telemóvel)
    t = wb.create_sheet("Top 5")
    t.append(["Categoria", "Prioridade", "Score", "Nome", "Instagram", "Canal", "Intuito", "Foco a investigar"])
    for c in t[1]:
        c.font, c.fill = cab, fundo
    for cat, rot in CATEGORIAS.items():
        do_cat = sorted([c for c in contas if c["categoria"] == cat], key=lambda c: -score_py(c)[0])[:5]
        for c in do_cat:
            s, p = score_py(c)
            t.append([rot, p, s, c["nome"], c.get("instagram", ""), M.canal(c), M.INTUITO[cat], c.get("foco", "")])
    for col, w in zip("ABCDEFGH", (22, 10, 7, 26, 20, 11, 50, 40)):
        t.column_dimensions[col].width = w

    # Comentários prontos para as comunidades
    k = wb.create_sheet("Comentários")
    k.append(["Comunidade", "Instagram", "Arranque 1", "Arranque 2", "Arranque 3"])
    for c in k[1]:
        c.font, c.fill = cab, fundo
    for c in contas:
        if c["categoria"] == "comunidade":
            k.append([c["nome"], c.get("instagram", "")] + M.comentarios(c))
    for col in "CDE":
        k.column_dimensions[col].width = 60

    # Regras do guia
    g = wb.create_sheet("Guia")
    linhas = [["Nota 0–10", "Encaixe 0–3 + Acesso 0–2 + Complexidade 0–3 + Prova 0–2"],
              ["Prioridade", "8–10 → 5 · 6–7 → 4 · 4–5 → 3 · 2–3 → 2 · 0–1 → 1 (descartar/observar)"],
              ["Envio", "1 minuto entre mensagens, saudação a variar; follow-up uma vez, 5 dias depois; depois «Sem resposta» e parar"],
              ["Linguagem", "Hipótese a validar, nunca acusação; não vender automação total; setores regulados sem decisões sensíveis"]]
    linhas += [["Estado", e] for e in ESTADOS]
    for l in linhas:
        g.append(l)
    g.column_dimensions["A"].width, g.column_dimensions["B"].width = 14, 100
    wb.save(saida)
    print("{} contas → {}".format(len(contas), saida))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
